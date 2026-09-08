"""정기점검 보고서 PPTX 생성.

회사에서 쓰던 양식(templates/monthly-report.pptx)을 그대로 열어 값만 갈아끼운다.
디자인·서식·마스터는 손대지 않는다.

  python scripts/build_report.py <입력.json> <출력.pptx>

입력 JSON 구조는 lib/pptx.ts 의 ReportPayload 와 같다.

왜 파이썬인가:
  PowerPoint 는 한 칸의 글을 여러 run 으로 잘게 쪼개 저장한다. 실제 양식에서
  "분석 및 진행 상황" 칸 하나가 48조각이었다. 문자열 치환으로는 못 찾는다.
  python-pptx 는 그 구조를 다루므로 첫 run 의 서식을 지키면서 안전하게 바꿀 수 있다.
"""
import copy
import json
import sys
from pathlib import Path

from pptx import Presentation

RELS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"

# 템플릿에서 각 구획이 시작하는 위치 (0-based)
IDX_CLOUD = 0        # 01 클라우드 운영 현황
IDX_LICENSE = 1      # 2-1 라이선스 현황 — TAS App Service 실 운영 현황 3칸만 갱신
IDX_SR_SUMMARY = 2   # 01 SR 진행현황 요약
IDX_SR_DETAIL = 3    # 2-2 SR 상세 (원본 1장)
IDX_WORK = 9         # 3-3 작업 진행 현황 (원본 1장)

WORK_ROWS_PER_SLIDE = 8   # 실제 양식이 한 장에 8행까지 담고 있다


# --------------------------------------------------------------------------
# 슬라이드 조작
# --------------------------------------------------------------------------

def clone_slide(prs, src):
    """슬라이드를 통째로 복제해 맨 뒤에 붙인다.

    python-pptx 에 공식 복제 기능이 없어 도형 XML 을 그대로 옮긴다.
    SR 상세·작업 슬라이드에는 그림이 없어(확인함) 관계(rel) 복사가 필요 없다.
    """
    dst = prs.slides.add_slide(src.slide_layout)
    for shape in list(dst.shapes):
        shape._element.getparent().remove(shape._element)
    for shape in src.shapes:
        dst.shapes._spTree.append(copy.deepcopy(shape._element))
    return dst


def slide_ids(prs):
    return list(prs.slides._sldIdLst)


def drop_slide(prs, element):
    """슬라이드 하나를 문서에서 뺀다."""
    rel_id = element.get(RELS + "id")
    prs.part.drop_rel(rel_id)
    prs.slides._sldIdLst.remove(element)


def reorder(prs, order):
    """order 에 담긴 sldId 요소 순서대로 다시 배열한다."""
    lst = prs.slides._sldIdLst
    for element in list(lst):
        lst.remove(element)
    for element in order:
        lst.append(element)


# --------------------------------------------------------------------------
# 텍스트 채우기
# --------------------------------------------------------------------------

def set_text(frame, text):
    """텍스트를 갈아끼우되 첫 run 의 서식(글꼴·크기·색)을 그대로 쓴다.

    한 칸이 수십 개 run 으로 쪼개져 있어도 첫 run 만 남기고 지운 뒤 채운다.
    여러 줄이면 첫 문단을 복제해 같은 서식을 유지한다.
    """
    text = "" if text is None else str(text)
    lines = text.split("\n")

    paragraphs = list(frame.paragraphs)
    first = paragraphs[0]

    # 첫 문단만 남긴다
    for extra in paragraphs[1:]:
        extra._p.getparent().remove(extra._p)

    def fill(paragraph, value):
        runs = list(paragraph.runs)
        if runs:
            runs[0].text = value
            for run in runs[1:]:
                run._r.getparent().remove(run._r)
        else:
            paragraph.text = value

    fill(first, lines[0])
    for line in lines[1:]:
        clone = copy.deepcopy(first._p)
        first._p.getparent().append(clone)
        # 복제한 문단의 run 을 새 값으로
        from pptx.text.text import _Paragraph
        fill(_Paragraph(clone, first._parent), line)


def cell_text(table, row, col, value):
    """표의 한 칸을 채운다. 범위를 벗어나면 조용히 넘어가지 않고 알린다."""
    if row >= len(table.rows) or col >= len(table.columns):
        raise IndexError(f"표 범위를 벗어남: ({row},{col}) / {len(table.rows)}x{len(table.columns)}")
    set_text(table.cell(row, col).text_frame, value)


def first_table(slide):
    for shape in slide.shapes:
        if shape.has_table:
            return shape.table
    raise LookupError("표가 없는 슬라이드")


def tables(slide):
    return [s.table for s in slide.shapes if s.has_table]


def replace_in_texts(slide, old, new):
    """슬라이드의 모든 텍스트에서 문자열을 바꾼다. 쪼개진 run 을 합쳐 비교한다."""
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        for paragraph in shape.text_frame.paragraphs:
            joined = "".join(r.text for r in paragraph.runs)
            if old in joined:
                runs = list(paragraph.runs)
                if runs:
                    runs[0].text = joined.replace(old, new)
                    for run in runs[1:]:
                        run._r.getparent().remove(run._r)


def set_page_no(slide, current, total):
    """제목의 "(1/6)" 같은 쪽 번호를 갱신한다."""
    import re
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        for paragraph in shape.text_frame.paragraphs:
            joined = "".join(r.text for r in paragraph.runs)
            if re.search(r"\(\s*\d+\s*/\s*\d+\s*\)", joined):
                fixed = re.sub(r"\(\s*\d+\s*/\s*\d+\s*\)", f"({current}/{total})", joined)
                runs = list(paragraph.runs)
                runs[0].text = fixed
                for run in runs[1:]:
                    run._r.getparent().remove(run._r)
                return


# --------------------------------------------------------------------------
# 표 행 늘리기 / 줄이기
# --------------------------------------------------------------------------

def resize_rows(table, keep_header, wanted):
    """데이터 행 수를 wanted 로 맞춘다.

    늘릴 때는 마지막 데이터 행을 복제한다 — 서식과 병합 상태를 그대로 물려받는다.
    """
    tbl = table._tbl
    rows = tbl.findall(
        "{http://schemas.openxmlformats.org/drawingml/2006/main}tr")
    have = len(rows) - keep_header
    if wanted == have:
        return
    if wanted < have:
        for element in rows[keep_header + wanted:]:
            tbl.remove(element)
        return
    template_row = rows[-1]
    for _ in range(wanted - have):
        tbl.append(copy.deepcopy(template_row))


# --------------------------------------------------------------------------
# 구획별 채우기
# --------------------------------------------------------------------------

def fill_cloud(slide, data, month_label):
    """01 클라우드 운영 현황."""
    table = first_table(slide)
    replace_in_texts(slide, table.cell(0, 2).text.strip(), f"{month_label} (누적)")

    # 행 2~7 = 여섯 구분, 행 8 = 합계
    for offset, row in enumerate(data["rows"]):
        r = 2 + offset
        cell_text(table, r, 4, row["container"])
        cell_text(table, r, 5, row["note"])
        cell_text(table, r, 8, row["delta"])

    total = data["total"]
    cell_text(table, 8, 2, total["cluster"])
    cell_text(table, 8, 3, total["host"])
    cell_text(table, 8, 4, total["container"])
    cell_text(table, 8, 8, total["delta"])


def fill_license(slide, values):
    """2-1 라이선스 현황 중 TAS App Service 행의 "실 운영 현황" 세 칸만 바꾼다.

    나머지 칸과 다른 제품 행은 양식 그대로 둔다.
    행 위치를 숫자로 박지 않고 제품명으로 찾는다 — 양식에 행이 추가돼도 버틴다.
    """
    for table in tables(slide):
        top = [table.cell(0, c).text.strip() for c in range(len(table.columns))]
        header = [table.cell(1, c).text.strip() for c in range(len(table.columns))]
        if "실 운영 현황" not in top:
            continue
        # "은행/중앙회/합계" 는 구독 계약 현황 아래에도 있다.
        # 반드시 "실 운영 현황" 이 시작하는 열부터 찾아야 구독 값을 덮어쓰지 않는다.
        start = top.index("실 운영 현황")
        try:
            bank = header.index("은행", start)
            central = header.index("중앙회", bank + 1)
            total = header.index("합계", central + 1)
        except ValueError:
            continue
        for r in range(len(table.rows)):
            names = " ".join(table.cell(r, c).text for c in range(3))
            if "TAS App Service" in names:
                cell_text(table, r, bank, values["bank"])
                cell_text(table, r, central, values["central"])
                cell_text(table, r, total, values["total"])
                return True
    return False


def fill_sr_summary(slide, items):
    """01 SR 진행현황 요약. 헤더 1행 + SR 수만큼."""
    table = first_table(slide)
    resize_rows(table, keep_header=1, wanted=len(items))

    for offset, sr in enumerate(items):
        r = 1 + offset
        cols = len(table.rows[r].cells)
        # 첫 행은 플랫폼·제품 칸이 살아 있고, 이어지는 행은 세로 병합으로 빠진다.
        base = cols - 5
        cell_text(table, r, base + 0, sr["no"])
        cell_text(table, r, base + 1, sr["openedOn"])
        cell_text(table, r, base + 2, sr["title"])
        cell_text(table, r, base + 3, sr["progress"])
        cell_text(table, r, base + 4, sr["done"])


def fill_sr_detail(slide, sr):
    """2-2 SR 상세 한 장."""
    table = first_table(slide)
    cell_text(table, 0, 1, sr["product"])
    cell_text(table, 0, 3, sr["no"])
    cell_text(table, 0, 5, sr["severity"])
    cell_text(table, 1, 1, sr["openedOn"])
    cell_text(table, 1, 3, sr["closedOn"])
    cell_text(table, 1, 5, sr["status"])
    cell_text(table, 2, 1, sr["title"])
    cell_text(table, 3, 1, sr["symptom"])
    cell_text(table, 4, 1, sr["analysis"])
    cell_text(table, 5, 1, sr["result"])


def fill_work(slide, rows):
    """3-3 작업 진행 현황 한 장."""
    table = first_table(slide)
    resize_rows(table, keep_header=1, wanted=len(rows))
    for offset, row in enumerate(rows):
        r = 1 + offset
        cell_text(table, r, 0, row["center"])
        cell_text(table, r, 1, row["corp"])
        cell_text(table, r, 2, row["span"])
        cell_text(table, r, 3, row["support"])
        cell_text(table, r, 4, row["title"])
        cell_text(table, r, 5, row["issue"])
        cell_text(table, r, 6, row["note"])


# --------------------------------------------------------------------------

def chunk(items, size):
    return [items[i:i + size] for i in range(0, len(items), size)] or [[]]


def build(payload, template_path, out_path):
    prs = Presentation(template_path)
    ids = slide_ids(prs)
    slides = list(prs.slides)

    sr_items = payload["srs"]
    work_pages = chunk(payload["work"], WORK_ROWS_PER_SLIDE)

    # 원본으로 쓸 슬라이드
    sr_proto = slides[IDX_SR_DETAIL]
    work_proto = slides[IDX_WORK]

    # 1) 고정 구획 채우기
    fill_cloud(slides[IDX_CLOUD], payload["cloud"], payload["monthLabel"])
    if payload.get("license"):
        if not fill_license(slides[IDX_LICENSE], payload["license"]):
            print("경고: 라이선스 표에서 TAS App Service 행을 찾지 못했습니다.", file=sys.stderr)
    fill_sr_summary(slides[IDX_SR_SUMMARY], sr_items)

    # 2) SR 상세 · 작업 슬라이드를 필요한 수만큼 새로 만든다
    made_sr = []
    for i, sr in enumerate(sr_items, 1):
        slide = clone_slide(prs, sr_proto)
        fill_sr_detail(slide, sr)
        set_page_no(slide, i, len(sr_items))
        made_sr.append(slide)

    made_work = []
    for i, rows in enumerate(work_pages, 1):
        slide = clone_slide(prs, work_proto)
        fill_work(slide, rows)
        set_page_no(slide, i, len(work_pages))
        made_work.append(slide)

    # 3) 원본 슬라이드들을 걷어내고 순서를 다시 잡는다
    new_ids = slide_ids(prs)
    made_ids = new_ids[len(ids):]
    keep_front = [ids[IDX_CLOUD], ids[IDX_LICENSE], ids[IDX_SR_SUMMARY]]
    made_sr_ids = made_ids[:len(made_sr)]
    made_work_ids = made_ids[len(made_sr):]

    for element in ids[IDX_SR_DETAIL:]:
        drop_slide(prs, element)

    reorder(prs, keep_front + made_sr_ids + made_work_ids)

    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    prs.save(out_path)
    return len(prs.slides._sldIdLst)


def main():
    if len(sys.argv) != 3:
        print("사용법: build_report.py <입력.json> <출력.pptx>", file=sys.stderr)
        return 1
    payload = json.loads(Path(sys.argv[1]).read_text(encoding="utf-8"))
    template = payload.get("template") or "templates/monthly-report.pptx"
    count = build(payload, template, sys.argv[2])
    print(json.dumps({"ok": True, "slides": count, "out": sys.argv[2]}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    sys.exit(main())
