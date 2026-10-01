# -*- coding: utf-8 -*-
"""만들어진 보고서를 재서 JSON 으로 찍는다. 읽기만 한다.

`scripts/build_report.py` 의 결과가 맞는지 확인하는 용도다. 사람이 열어 보고
"사진이 좀 벌어진 것 같다" 를 판단하는 대신 숫자를 본다 — 실제로 눈으로만 보다가
가운데가 1인치 벌어진 보고서와 지난달 바닥글이 달린 보고서를 배포했다.

    .venv/Scripts/python.exe scripts/describe_report.py 보고서.pptx

test/reportBuildPptx.test.ts 가 이걸 불러 결과를 검사한다. 출력 모양을 바꾸면
그 시험도 같이 고쳐야 한다.
"""
from __future__ import annotations

import json
import re
import sys

from pptx import Presentation

EMU_PER_INCH = 914400.0
PICTURE = 13       # MSO_SHAPE_TYPE.PICTURE
PLACEHOLDER = 14   # MSO_SHAPE_TYPE.PLACEHOLDER

# 달·기준일이 박혀 있던 자리를 찾는 모양. build_report.py 의 두 패턴과 짝이다.
LABEL_PATTERN = re.compile(r"\d{4}\s*년\s*\d{1,2}\s*월|\d{4}\.\d{2}\.\d{2}\s*기준")


def inches(value: int) -> float:
    return round(value / EMU_PER_INCH, 3)


def labels_of(slide) -> list[str]:
    """이 장에 보이는 달·기준일 문구. 레이아웃에 있는 것까지 포함한다.

    바닥글은 슬라이드가 아니라 **레이아웃**에 있다. 슬라이드만 보면 못 찾는다 —
    그래서 지난달 라벨이 12장에 그대로 붙어 나갔다.
    """
    found = set()
    for shapes in (slide.shapes, slide.slide_layout.shapes):
        for shape in shapes:
            if not shape.has_text_frame:
                continue
            text = shape.text_frame.text.strip()
            if LABEL_PATTERN.search(text):
                found.add(text)
    return sorted(found)


# 구간 제목("1. Summary" 처럼 번호로 시작)과 머리말 칸을 가르는 기준.
SECTION_PATTERN = re.compile(r"^\d+\.\s")
CHIP_MAX_WIDTH = 914400   # 1.0in. build_report.py 의 CHIP_MAX_WIDTH 와 같은 기준


def section_of(slide) -> str:
    """이 장의 구간 제목. **레이아웃에서** 온다.

    고객 문서는 구간을 레이아웃으로 구분한다 — 클라우드·라이선스는 '1. Summary',
    SR·작업은 '2. 이슈 및 작업 진행 사항', 사진은 '5. 정기점검 결과 및 상세 내역'.
    한때 우리 양식이 레이아웃 하나뿐이어서 모든 장이 '2. …' 로 나갔다.
    """
    for shape in slide.slide_layout.shapes:
        if shape.has_text_frame and SECTION_PATTERN.match(shape.text_frame.text.strip()):
            return shape.text_frame.text.strip()
    return ""


def header_of(slide) -> dict:
    """머리말의 구획 번호와 제목. 폭으로 가른다(좁은 쪽이 구획 번호)."""
    chip = title = ""
    for shape in slide.shapes:
        if not shape.has_text_frame or shape.width is None:
            continue
        if shape.shape_type != PLACEHOLDER:
            continue
        if shape.width <= CHIP_MAX_WIDTH:
            chip = shape.text_frame.text.strip()
        else:
            title = shape.text_frame.text.strip()
    return {"chip": chip, "title": title}


def footer_font(prs) -> dict:
    """바닥글 연·월 run 의 글꼴과 크기. 레이아웃마다 다를 수 있어 모아서 돌려준다."""
    found = set()
    for master in prs.slide_masters:
        for layout in master.slide_layouts:
            for shape in layout.shapes:
                if not shape.has_text_frame:
                    continue
                if not LABEL_PATTERN.search(shape.text_frame.text):
                    continue
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        found.add((run.font.name,
                                   run.font.size.pt if run.font.size else None))
    return {"names": sorted({n or "" for n, _ in found}),
            "sizes": sorted({s for _, s in found if s is not None})}


def tables_of(slide) -> list:
    """표의 자리와 열 폭·행 높이. 장마다 미세하게 어긋나던 값들이다."""
    out = []
    for shape in slide.shapes:
        if not shape.has_table:
            continue
        table = shape.table
        out.append({
            "left": inches(shape.left),
            "top": inches(shape.top),
            "width": inches(shape.width),
            "columns": [inches(c.width) for c in table.columns],
            "rows": [inches(r.height) for r in table.rows],
        })
    return out


def describe(path: str) -> dict:
    prs = Presentation(path)
    slides = []
    for slide in prs.slides:
        photos = [
            {
                "left": inches(shape.left),
                "top": inches(shape.top),
                "width": inches(shape.width),
                "height": inches(shape.height),
            }
            for shape in slide.shapes
            if shape.shape_type == PICTURE
        ]
        slides.append({
            "layout": slide.slide_layout.name,
            "section": section_of(slide),
            "header": header_of(slide),
            "photos": photos,
            "labels": labels_of(slide),
            "tables": tables_of(slide),
        })
    return {
        "slideWidth": inches(prs.slide_width),
        "slideHeight": inches(prs.slide_height),
        "footerFont": footer_font(prs),
        "slides": slides,
    }


def main(argv: list[str]) -> int:
    if len(argv) != 2:
        print("사용법: describe_report.py <보고서.pptx>", file=sys.stderr)
        return 1
    # Windows 기본 출력 인코딩이 cp949 라 한글 라벨이 깨져 나간다. 부르는 쪽은 UTF-8
    # 로 읽으므로 글자가 뭉개져도 **오류 없이** 엉뚱한 문자열이 전달된다 — 실제로
    # 시험이 "라벨을 못 찾았다" 로 헛돌았다. 여기서 UTF-8 로 못 박는다.
    sys.stdout.reconfigure(encoding="utf-8")
    json.dump(describe(argv[1]), sys.stdout, ensure_ascii=False)
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
